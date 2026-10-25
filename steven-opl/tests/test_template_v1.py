#!/usr/bin/env python3
import io,json,subprocess,sys,tempfile,zipfile
from pathlib import Path
from PIL import Image
from docx import Document
ROOT=Path(__file__).resolve().parents[1]; PY=sys.executable
TEMPLATE=ROOT/'assets/templates/steven-opl-standard/1.1.0/template.docx'; CONTRACT=TEMPLATE.with_name('contract.json')

def fact(v,status='user_stated',refs=None): return {'value':v,'status':status,'source_refs':refs if refs is not None else ['MSG-001']}
def run(*args,expect=0):
 p=subprocess.run([PY,*map(str,args)],capture_output=True,text=True)
 if p.returncode!=expect: raise AssertionError(f"{args}\n{p.stdout}\n{p.stderr}")
 return p

def sha(path):
 import hashlib
 return hashlib.sha256(Path(path).read_bytes()).hexdigest()

def main():
 with tempfile.TemporaryDirectory() as td:
  td=Path(td); bmp, png, exifjpg=td/'a.bmp',td/'b.png',td/'c.jpg'
  Image.new('RGB',(120,80),'red').save(bmp)
  Image.new('RGBA',(60,120),(0,100,255,180)).save(png)
  im=Image.new('RGB',(40,80),'green'); exif=im.getexif(); exif[274]=6; im.save(exifjpg,exif=exif)
  paths=[bmp,png,exifjpg]; image_ids=['IMG-001','IMG-002','IMG-003']; asset_ids=['ASSET-001','ASSET-002','ASSET-003']
  sources=[{'source_id':'MSG-001','type':'message'},{'source_id':'RULE:lesson-type-map-v1','type':'rule'}]
  sources += [{'source_id':a,'type':'image','path':str(p),'sha256':sha(p)} for a,p in zip(asset_ids,paths)]
  images=[]
  for i,(iid,aid,p) in enumerate(zip(image_ids,asset_ids,paths),1):
   images.append({'image_id':iid,'asset_id':aid,'path':str(p),'source_sha256':sha(p),'display_width_px':120,'display_height_px':80,
    'original_format':p.suffix[1:].upper(),'send_order':i,'work_order':i,'insert_order':i,'user_description':'测试图','usage':'测试要点',
    'step_id':'STEP-001' if i<3 else 'STEP-003','caption':f'图{i}','included':True,'confirmed':True,'source_refs':[aid,'MSG-001'],
    'transform_authorization':{'rotate':False,'mirror':False,'crop':False}})
  draft={'schema_version':'1.1','task_id':'T','draft_id':'D','lesson_type':'repair_tip','source_index':sources,
   'scope':{'single_point':True,'topic':fact('测试单点'),'split_required':False,'excluded_topics':[]},
   'metadata':{'group':fact('维修组'),'procedure':fact('测试工序'),'number':fact('T-001'),'creation_date':fact('2026-10-26'),
    'template_category':fact('基础知识','derived',['RULE:lesson-type-map-v1']),'director':fact(None,'unknown',[]),
    'team_leader':fact(None,'unknown',[]),'made_by':fact('测试员')},
   'content':{k:fact(v) for k,v in {'title':'测试一点课','applicable_to':'测试设备','prerequisites':'断电','problem_or_goal':'测试问题',
    'key_point':'测试要点','mistakes':'测试易错点','safety':'先断电','acceptance':'检查合格','limitations':'测试限制'}.items()},
   'steps':[{'step_id':f'STEP-{i:03d}','order':i,'instruction':fact(f'操作{i}'),'check':fact(f'检查{i}'),'source_refs':['MSG-001']} for i in range(1,4)],
   'images':images,'open_items':[],'template_selection':None}
  dp,approved=td/'draft.json',td/'approved.json'; dp.write_text(json.dumps(draft,ensure_ascii=False),encoding='utf-8')
  run(ROOT/'scripts/approve.py','--draft',dp,'--output',approved,'--source-ref','MSG-APPROVE-CONTENT')
  ad=json.loads(approved.read_text('utf-8'))
  from hashlib import sha256
  canonical=lambda x:sha256(json.dumps(x,ensure_ascii=False,sort_keys=True,separators=(',',':')).encode()).hexdigest()
  plan={'schema_version':'1.0','task_id':'T','plan_id':'P-1','approved_sha256':canonical(ad),
   'template_selection':{'template_id':'steven-opl-standard','template_version':'1.1.0'},
   'rows':[{'row_id':'ROW-001','title':'双图要点','text':'两张图片服务同一个教学点','image_ids':['IMG-001','IMG-002'],'content_refs':['STEP-001'],'image_side':'left'},
           {'row_id':'ROW-002','title':'无图要点','text':'无图时使用整行文字','image_ids':[],'content_refs':['STEP-002'],'image_side':'full_width'},
           {'row_id':'ROW-003','title':'方向要点','text':'保持手机照片原始显示方向','image_ids':['IMG-003'],'content_refs':['STEP-003'],'image_side':'left'}]}
  pp,pap=td/'plan.json',td/'plan-approved.json'; pp.write_text(json.dumps(plan,ensure_ascii=False),encoding='utf-8')
  run(ROOT/'scripts/approve_render_plan.py','--plan',pp,'--approved',approved,'--output',pap,'--source-ref','MSG-APPROVE-LAYOUT')
  out,trace,report=td/'out.docx',td/'trace.json',td/'report.json'
  run(ROOT/'scripts/fill_word_opl.py','--approved',approved,'--plan',pap,'--template',TEMPLATE,'--contract',CONTRACT,'--output',out,'--trace',trace)
  run(ROOT/'scripts/validate_output.py','--approved',approved,'--plan',pap,'--contract',CONTRACT,'--docx',out,'--trace',trace,'--report',report)
  result=json.loads(report.read_text('utf-8')); assert result['delivery_status']=='ready_pending_visual_review'
  tr=json.loads(trace.read_text('utf-8')); target=next(x for x in tr['images'] if x['image_id']=='IMG-003')['prepared_sha256']
  import hashlib
  with zipfile.ZipFile(out) as z:
   payload=next(z.read(n) for n in z.namelist() if n.startswith('word/media/') and hashlib.sha256(z.read(n)).hexdigest()==target)
  with Image.open(io.BytesIO(payload)) as displayed: assert displayed.width>displayed.height, displayed.size
  tampered=td/'training-filled.docx'; doc=Document(out); doc.tables[1].cell(0,5).text='2026-10-26'; doc.save(tampered)
  run(ROOT/'scripts/validate_template_layout.py',tampered,'--plan',pap,'--json',expect=1)
  Image.new('RGB',(10,10),'black').save(bmp)
  run(ROOT/'scripts/validate_render_plan.py','--plan',pap,'--approved',approved,expect=1)
  print('dynamic template workflow passed')
if __name__=='__main__': main()

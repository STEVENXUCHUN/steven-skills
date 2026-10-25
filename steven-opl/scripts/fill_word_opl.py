#!/usr/bin/env python3
"""Render approved Steven OPL content + approved render plan into DOCX."""
import argparse, os, shutil, subprocess, sys, tempfile
from copy import deepcopy
from pathlib import Path
from common import dump_json, file_sha256, get_path, load_json

try:
    from PIL import Image, ImageOps
    from docx import Document
    from docx.enum.table import WD_ALIGN_VERTICAL, WD_ROW_HEIGHT_RULE
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt
except ImportError as exc:
    raise SystemExit(f"required document dependency unavailable: {exc}")


def set_run(run, text, style):
    run.text = str(text or "")
    run.font.size = Pt(style.get("size_pt", 11)); run.font.bold = style.get("bold", False)
    run.font.name = style.get("font_en", "Times New Roman")
    rpr = run._element.get_or_add_rPr(); fonts = rpr.find(qn("w:rFonts"))
    if fonts is None: fonts = OxmlElement("w:rFonts"); rpr.append(fonts)
    fonts.set(qn("w:eastAsia"), style.get("font_cn", "宋体"))


def clear_cell(cell):
    for p in cell.paragraphs[1:]: p._element.getparent().remove(p._element)
    p = cell.paragraphs[0]
    for r in list(p.runs): r._element.getparent().remove(r._element)
    return p


def replace_marker(doc, marker, value):
    hits = 0
    for node in doc.part.element.iter(qn("w:t")):
        if node.text and marker in node.text:
            node.text = node.text.replace(marker, str(value or "")); hits += 1
    return hits


def find_bookmark(doc, name):
    for el in doc.part.element.iter(qn("w:bookmarkStart")):
        if el.get(qn("w:name")) == name: return el
    return None


def write_bookmark_text(doc, name, value):
    mark = find_bookmark(doc, name)
    if mark is None: return False
    end = None
    for el in mark.getparent().iter(qn("w:bookmarkEnd")):
        if el.get(qn("w:id")) == mark.get(qn("w:id")): end = el; break
    if end is None: return False
    run = OxmlElement("w:r"); text = OxmlElement("w:t"); text.text = str(value or ""); run.append(text); end.addprevious(run)
    return True


def set_cant_split(row):
    trpr = row._tr.get_or_add_trPr()
    if trpr.find(qn("w:cantSplit")) is None: trpr.append(OxmlElement("w:cantSplit"))


def ensure_rows(table, count):
    if not table.rows: raise ValueError("lesson table has no template row")
    templates = [deepcopy(table.rows[0]._tr), deepcopy(table.rows[min(1, len(table.rows)-1)]._tr)]
    while len(table.rows) < count: table._tbl.append(deepcopy(templates[len(table.rows) % 2]))
    while len(table.rows) > count:
        tr = table.rows[count]._tr; tr.getparent().remove(tr)


def prepare_image(item, tmpdir, quality):
    source = item["path"]
    with Image.open(source) as opened:
        image = ImageOps.exif_transpose(opened)
        image.thumbnail((2400, 2400) if quality == "high" else (1800, 1800), Image.Resampling.LANCZOS)
        keep_png = (opened.format == "PNG") or ("A" in image.getbands())
        ext = ".png" if keep_png else ".jpg"
        out = Path(tmpdir) / f"{item['image_id']}{ext}"
        if keep_png:
            image.save(out, "PNG", optimize=True)
        else:
            image.convert("RGB").save(out, "JPEG", quality=92 if quality == "high" else 86, optimize=True)
    return out


def fitted_dimension(path, max_width, max_height):
    with Image.open(path) as image: width, height = image.size
    if width / height >= max_width / max_height: return {"width": Cm(max_width)}
    return {"height": Cm(max_height)}


def insert_images(cell, items, tmpdir, box, quality, trace, row_no, side):
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    first = clear_cell(cell); first.alignment = WD_ALIGN_PARAGRAPH.CENTER
    prepared = []
    for item in items:
        path = prepare_image(item, tmpdir, quality); prepared.append((item, path))
        trace.append({"image_id": item["image_id"], "row": row_no, "side": side,
                      "source_sha256": item["source_sha256"], "prepared_sha256": file_sha256(path),
                      "prepared_format": path.suffix.lower()})
    pairs = [prepared[i:i+2] for i in range(0, len(prepared), 2)]
    for pi, pair in enumerate(pairs):
        p = first if pi == 0 else cell.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(2)
        for item, path in pair:
            if len(prepared) == 1:
                dims = fitted_dimension(path, box["single_width"], box["single_height"])
            else:
                dims = fitted_dimension(path, box["multi_width"], box["multi_height"])
            p.add_run().add_picture(str(path), **dims)


def fill_row_text(cell, title, text, styles):
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    title_p = clear_cell(cell); title_p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    title_p.paragraph_format.space_before = Pt(4); title_p.paragraph_format.space_after = Pt(6)
    set_run(title_p.add_run(), title, styles["title"])
    body = cell.add_paragraph(); body.alignment = WD_ALIGN_PARAGRAPH.LEFT
    body.paragraph_format.line_spacing = styles["body"].get("line_spacing", 1.3); body.paragraph_format.space_after = Pt(4)
    set_run(body.add_run(), text, styles["body"])


def fill_header(doc, approved, contract):
    for name, slot in contract.get("text_slots", {}).items():
        try: value = get_path(approved, slot["source"])
        except Exception: value = None
        if slot.get("required") and (value is None or not str(value).strip()):
            raise ValueError(f"required field is empty: {name}")
        if slot.get("bookmark"):
            found = write_bookmark_text(doc, slot["bookmark"], value)
            if slot.get("required") and not found: raise ValueError(f"required bookmark missing: {name}")
        else:
            hits = replace_marker(doc, slot.get("marker", "{{" + name + "}}"), value)
            if slot.get("required") and hits == 0: raise ValueError(f"required marker missing: {name}")


def fill_lesson_table(doc, approved, plan, contract, tmpdir, quality, trace):
    cfg = contract["dynamic_lesson_table"]
    main = doc.tables[cfg["main_table_index"]]
    tables = main.cell(cfg["container_row"], cfg["container_col"]).tables
    if len(tables) != 1: raise ValueError("template lesson table not found")
    table, rows = tables[0], plan["rows"]
    if not cfg["min_rows"] <= len(rows) <= cfg["max_rows"]: raise ValueError("render row count outside template contract")
    ensure_rows(table, len(rows))
    images = {x["image_id"]: x for x in approved.get("images", [])}
    styles = {"title": cfg["title_style"], "body": cfg["body_style"]}; box = cfg["image_box_cm"]
    for index, spec in enumerate(rows):
        row = table.rows[index]; set_cant_split(row); row.height = Cm(cfg["row_height_cm"]); row.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
        left, right = row.cells[:2]; items = [images[x] for x in spec["image_ids"]]
        if not items:
            fill_row_text(left.merge(right), spec["title"], spec["text"], styles)
            continue
        image_cell, text_cell = (left, right) if index % 2 == 0 else (right, left)
        side = "left" if index % 2 == 0 else "right"
        insert_images(image_cell, items, tmpdir, box, quality, trace, index + 1, side)
        fill_row_text(text_cell, spec["title"], spec["text"], styles)
    return table


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--approved", required=True); ap.add_argument("--plan", required=True)
    ap.add_argument("--template", required=True); ap.add_argument("--contract", required=True)
    ap.add_argument("--output", required=True); ap.add_argument("--trace", required=True)
    ap.add_argument("--image-quality", choices=("standard", "high"), default="standard")
    args = ap.parse_args()
    approved, plan, contract = load_json(args.approved), load_json(args.plan), load_json(args.contract)
    scripts = Path(__file__).parent
    if subprocess.run([sys.executable, str(scripts / "validate_data.py"), args.approved]).returncode:
        raise SystemExit("approved content validation failed")
    if subprocess.run([sys.executable, str(scripts / "validate_render_plan.py"), "--plan", args.plan, "--approved", args.approved]).returncode:
        raise SystemExit("approved render-plan validation failed")
    if approved.get("schema_version") != contract.get("supported_schema"):
        raise SystemExit("content schema is not supported by template")
    if plan.get("template_selection", {}) != {"template_id": contract.get("template_id"), "template_version": contract.get("template_version"}:
        raise SystemExit("render plan template selection does not match contract")
    template_hash = file_sha256(args.template)
    if template_hash != contract.get("template_sha256"): raise SystemExit("template SHA-256 mismatch")
    out = Path(args.output)
    if Path(args.template).resolve() == out.resolve(): raise SystemExit("refusing to overwrite template")
    out.parent.mkdir(parents=True, exist_ok=True); Path(args.trace).parent.mkdir(parents=True, exist_ok=True)
    fd, tmpname = tempfile.mkstemp(suffix=".docx", dir=str(out.parent)); os.close(fd)
    image_tmp = tempfile.mkdtemp(prefix="steven_opl_images_")
    image_trace = []
    try:
        shutil.copyfile(args.template, tmpname); os.chmod(tmpname, 0o600)
        doc = Document(tmpname)
        training_before = doc.tables[1]._tbl.xml if len(doc.tables) > 1 else None
        fill_header(doc, approved, contract)
        lesson = fill_lesson_table(doc, approved, plan, contract, image_tmp, args.image_quality, image_trace)
        training_after = doc.tables[1]._tbl.xml if len(doc.tables) > 1 else None
        if contract.get("training_record_policy") == "leave_blank" and training_before != training_after:
            raise ValueError("training record table changed")
        doc.save(tmpname)
        if file_sha256(args.template) != template_hash: raise ValueError("template source changed during rendering")
        os.replace(tmpname, out)
        trace = {
            "template_id": contract.get("template_id"), "template_version": contract.get("template_version"),
            "template_sha256": template_hash, "output": str(out), "output_sha256": file_sha256(out),
            "row_count": len(plan["rows"]), "training_record_untouched": training_before == training_after,
            "images": image_trace
        }
        dump_json(trace, args.trace)
        print(f"wrote candidate: {out}; rows={len(plan['rows'])}; images={len(image_trace)}")
    except Exception as exc:
        print(f"generation failed: {exc}", file=sys.stderr); raise SystemExit(1)
    finally:
        if os.path.exists(tmpname): os.unlink(tmpname)
        shutil.rmtree(image_tmp, ignore_errors=True)

if __name__ == "__main__": main()

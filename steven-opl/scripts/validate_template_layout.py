#!/usr/bin/env python3
"""Validate steven-opl-standard 1.1.0 dynamic DOCX layout."""
import argparse, json, sys, zipfile
from pathlib import Path
from docx import Document

EXPECTED_PAGE={"width":21.0,"height":29.7,"left":2.5,"right":0.9,"top":1.37,"bottom":0.55}
TRAINING_LABELS={"培训记录TrainingRecord","日期Date","讲师Trainer","受培训者Trainee"}

def compact(value): return "".join(str(value or "").split())
def has_cant_split(row): return bool(row._tr.xpath('./w:trPr/w:cantSplit'))
def image_count(cell): return len(cell._tc.xpath('.//wp:inline'))

def check(path, plan):
    out={"file":str(path),"passed":False,"errors":[],"warnings":[],"rows":[],"training_record":"unchecked"}
    try:
        with zipfile.ZipFile(path) as z:
            bad=z.testzip()
            if bad: out["errors"].append(f"DOCX损坏: {bad}")
    except Exception as exc:
        out["errors"].append(f"无效DOCX: {exc}"); return out
    doc=Document(path); section=doc.sections[0]
    actual={"width":section.page_width.cm,"height":section.page_height.cm,"left":section.left_margin.cm,
            "right":section.right_margin.cm,"top":section.top_margin.cm,"bottom":section.bottom_margin.cm}
    for key,value in EXPECTED_PAGE.items():
        if abs(actual[key]-value)>0.08: out["errors"].append(f"页面参数{key}偏离模板: {actual[key]:.2f}")
    if len(doc.tables)!=2:
        out["errors"].append(f"顶层表格应为2个，实际{len(doc.tables)}"); return out
    main,training=doc.tables
    if (len(main.rows),len(main.columns))!=(6,7): out["errors"].append("主表不是6x7")
    if (len(training.rows),len(training.columns))!=(4,15): out["errors"].append("培训表不是4x15")
    unexpected=[]
    for row in training.rows:
        for cell in row.cells:
            text=compact(cell.text)
            if text and text not in TRAINING_LABELS: unexpected.append(text)
    if unexpected:
        out["errors"].append("培训记录必须留空，发现填写内容: "+", ".join(sorted(set(unexpected))))
        out["training_record"]="not_blank"
    else: out["training_record"]="blank"
    nested=main.cell(5,0).tables
    if len(nested)!=1:
        out["errors"].append(f"主体嵌套表应为1个，实际{len(nested)}"); return out
    table=nested[0]; specs=plan.get("rows",[])
    if len(table.rows)!=len(specs): out["errors"].append(f"教学行数应为{len(specs)}，实际{len(table.rows)}")
    if len(table.columns)!=2: out["errors"].append("教学表不是2列")
    for i,row in enumerate(table.rows):
        if i>=len(specs): break
        spec=specs[i]; ids=spec.get("image_ids",[]); left,right=row.cells[:2]
        counts=[image_count(left),image_count(right)]; merged=left._tc is right._tc
        expected_side=0 if i%2==0 else 1; other=1-expected_side
        if ids:
            if merged: out["errors"].append(f"第{i+1}行有图但单元格被合并")
            if counts[expected_side]!=len(ids): out["errors"].append(f"第{i+1}行图片数应为{len(ids)}，实际{counts[expected_side]}")
            if counts[other]: out["errors"].append(f"第{i+1}行图片出现在错误侧")
            text_cell=right if expected_side==0 else left
        else:
            if not merged: out["errors"].append(f"第{i+1}行无图时应合并为整行文字")
            if sum(counts): out["errors"].append(f"第{i+1}行不应有图片")
            text_cell=left
        text=compact(text_cell.text)
        if compact(spec.get("title")) not in text or compact(spec.get("text")) not in text:
            out["errors"].append(f"第{i+1}行文字与批准排版计划不一致")
        if not has_cant_split(row): out["errors"].append(f"第{i+1}行缺少cantSplit")
        if row.height and row.height.cm<5.5: out["errors"].append(f"第{i+1}行高度不足")
        out["rows"].append({"row":i+1,"expected_side":"full_width" if not ids else ("left" if expected_side==0 else "right"),
                            "expected_images":len(ids),"image_counts":counts,"merged":merged,"cant_split":has_cant_split(row)})
    if "{{" in "\n".join(p.text for p in doc.paragraphs) or any("{{" in c.text for t in doc.tables for r in t.rows for c in r.cells):
        out["errors"].append("文档仍含占位符")
    out["passed"]=not out["errors"]; return out

def main():
    ap=argparse.ArgumentParser(); ap.add_argument('docx',type=Path); ap.add_argument('--plan',required=True); ap.add_argument('--json',action='store_true')
    args=ap.parse_args(); plan=json.loads(Path(args.plan).read_text('utf-8')); report=check(args.docx,plan)
    print(json.dumps(report,ensure_ascii=False,indent=2) if args.json else ('PASS' if report['passed'] else 'FAIL'))
    if not args.json:
        for item in report['errors']: print('ERROR:',item)
    return 0 if report['passed'] else 1
if __name__=='__main__': raise SystemExit(main())
